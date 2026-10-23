"""Standalone courtroom document endpoint tests, without external services."""

from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PyPDF2 import PdfWriter
from PyPDF2.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest

from src.features.courtroom.document_router import MAX_FILE_BYTES, MAX_TEXT_CHARS, router


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(router)
    with TestClient(app) as test_client:
        yield test_client


def upload(client, filename, content):
    return client.post(
        "/api/v1/courtroom/document", files={"file": (filename, content)}
    )


def make_pdf(*pages, password=None):
    writer = PdfWriter()
    for text in pages:
        writer.add_blank_page(width=300, height=200)
        page = writer.pages[-1]
        if text:
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
            })
            content = DecodedStreamObject()
            content.set_data(f"BT /F1 12 Tf 30 100 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = content
    if password:
        writer.encrypt(password)
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


def test_extract_txt(client):
    response = upload(client, "facts.TXT", b"\xef\xbb\xbf  Legal facts\nSecond line  ")
    assert response.status_code == 200
    assert response.json() == {
        "text": "Legal facts\nSecond line", "filename": "facts.TXT",
        "page_count": None, "extracted_characters": 23,
        "truncated": False, "warnings": [],
    }


def test_extract_docx(client):
    document = Document()
    document.add_paragraph("Petitioner facts")
    document.add_paragraph("Respondent facts")
    stream = BytesIO()
    document.save(stream)
    response = upload(client, "facts.docx", stream.getvalue())
    assert response.status_code == 200
    assert response.json() == {
        "text": "Petitioner facts\nRespondent facts", "filename": "facts.docx",
        "page_count": None, "extracted_characters": 33,
        "truncated": False, "warnings": [],
    }


def test_extract_pdf(client):
    response = upload(client, "facts.pdf", make_pdf("Legal facts"))
    assert response.status_code == 200
    assert response.json() == {
        "text": "[PDF page 1]\nLegal facts", "filename": "facts.pdf",
        "page_count": 1, "extracted_characters": 24,
        "truncated": False, "warnings": [],
    }


def test_pdf_page_markers(client):
    response = upload(client, "facts.pdf", make_pdf("First page", "Second page"))
    assert response.status_code == 200
    result = response.json()
    assert result["page_count"] == 2
    assert result["text"] == "[PDF page 1]\nFirst page\n[PDF page 2]\nSecond page"
    assert result["extracted_characters"] == len(result["text"])
    assert result["warnings"] == []


def test_pdf_partial_empty_page(client):
    response = upload(client, "facts.pdf", make_pdf("Legal facts", "", "More facts"))
    assert response.status_code == 200
    result = response.json()
    assert result["page_count"] == 3
    assert "[PDF page 2]" in result["text"]
    assert "[PDF page 3]\nMore facts" in result["text"]
    assert len(result["warnings"]) == 1
    assert "page 2 has no extractable text" in result["warnings"][0]
    assert "OCR is not performed" in result["warnings"][0]


def test_encrypted_pdf(client):
    response = upload(client, "facts.pdf", make_pdf("Legal facts", password="secret"))
    assert response.status_code == 422
    assert "encrypted" in response.json()["detail"]
    assert "password" in response.json()["detail"]


def test_docx_table_in_reading_order(client):
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Party", "Claim"), ("Petitioner", "Damages")]):
        for cell, text in zip(row.cells, values):
            cell.text = text
    document.add_paragraph("After table")
    stream = BytesIO()
    document.save(stream)
    response = upload(client, "facts.docx", stream.getvalue())
    assert response.status_code == 200
    assert response.json()["text"] == (
        "Before table\nParty\tClaim\nPetitioner\tDamages\nAfter table"
    )


@pytest.mark.parametrize("extension", ["txt", "docx", "pdf"])
def test_text_is_bounded(client, extension):
    content = b"a" * (MAX_TEXT_CHARS + 100)
    if extension == "docx":
        document = Document()
        document.add_paragraph(content.decode())
        stream = BytesIO()
        document.save(stream)
        content = stream.getvalue()
    elif extension == "pdf":
        content = make_pdf(content.decode(), "", "Later text")
    response = upload(client, f"facts.{extension}", content)
    assert response.status_code == 200
    result = response.json()
    assert len(result["text"]) == MAX_TEXT_CHARS
    assert result["extracted_characters"] == MAX_TEXT_CHARS
    assert result["truncated"] is True
    assert any("truncated" in warning for warning in result["warnings"])
    if extension == "pdf":
        assert result["text"].startswith("[PDF page 1]\n")
        assert result["page_count"] == 3
        assert any("page 2 has no extractable text" in warning
                   for warning in result["warnings"])
    else:
        assert result["text"] == "a" * MAX_TEXT_CHARS


def test_exact_text_limit_not_truncated(client):
    response = upload(client, "facts.txt", b"a" * MAX_TEXT_CHARS)
    assert response.status_code == 200
    assert response.json()["truncated"] is False
    assert response.json()["warnings"] == []


def test_oversized_file(client):
    response = upload(client, "facts.txt", b"a" * (MAX_FILE_BYTES + 1))
    assert response.status_code == 413


def test_exact_size_limit(client):
    response = upload(client, "facts.txt", b"a" * MAX_FILE_BYTES)
    assert response.status_code == 200
    assert len(response.json()["text"]) == MAX_TEXT_CHARS


@pytest.mark.parametrize("filename", ["facts.exe", "facts", "facts.doc"])
def test_unsupported_file(client, filename):
    assert upload(client, filename, b"facts").status_code == 400


@pytest.mark.parametrize(
    "filename,content",
    [
        ("facts.txt", b""),
        ("facts.txt", b" \r\n\t"),
        ("facts.txt", b"\xff\xfe"),
        ("facts.txt", b"binary\x00content"),
        ("facts.pdf", b"not a PDF"),
        ("facts.docx", b"not a DOCX"),
    ],
)
def test_unreadable_or_blank_document(client, filename, content):
    assert upload(client, filename, content).status_code == 422


def test_blank_docx(client):
    stream = BytesIO()
    Document().save(stream)
    assert upload(client, "blank.docx", stream.getvalue()).status_code == 422


def test_blank_pdf(client):
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(stream)
    assert upload(client, "blank.pdf", stream.getvalue()).status_code == 422


def test_missing_file(client):
    assert client.post("/api/v1/courtroom/document").status_code == 422


def test_downloadable_indian_sample_is_readable_case_input(client):
    sample = Path(__file__).resolve().parents[1] / "frontend/public/samples/indian-theft-case.pdf"
    response = upload(client, sample.name, sample.read_bytes())
    assert response.status_code == 200
    result = response.json()
    assert result["page_count"] == 3
    assert not result["truncated"] and result["warnings"] == []
    assert "[PDF page 3]" in result["text"]
    assert "Maharashtra" in result["text"]
    assert "[E1]" in result["text"] and "[E7]" in result["text"]
    assert "FICTIONAL" in result["text"]
    assert 60 < result["extracted_characters"] < 25000
